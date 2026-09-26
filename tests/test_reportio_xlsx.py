"""Tests for :mod:`vinyllabels.reportio.xlsx`."""

from __future__ import annotations

from collections.abc import Callable, Sequence
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from vinyllabels.reportio import xlsx as xlsx_module
from vinyllabels.reportio.xlsx import (
    Sheet,
    XlsxColumn,
    sanitize_sheet_name,
    select_xlsx_columns,
    write_workbook,
)


class _NoDefaultSheetWorkbook(Workbook):
    """A workbook born without openpyxl's default sheet, so ``active`` is None."""

    def __init__(self) -> None:
        super().__init__()
        self.remove(self["Sheet"])


def columns() -> list[XlsxColumn[tuple[str, float]]]:
    return [
        XlsxColumn("Name", lambda row: row[0]),
        XlsxColumn("Amount", lambda row: row[1], "amount"),
        XlsxColumn("Hidden", lambda row: "x", "hidden"),
    ]


def rows() -> list[tuple[str, float]]:
    return [("alpha", 1.5), ("b", 2.0)]


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("plain", "plain"),
        ("a/b\\c?d*e:f[g]", "a-b-c-d-e-f-g-"),
        ("x" * 40, "x" * 31),
    ],
)
def test_sanitize_sheet_name(name: str, expected: str) -> None:
    assert sanitize_sheet_name(name) == expected


def test_write_workbook_without_a_default_sheet(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Nothing to remove when the workbook has no active sheet; the writer must
    # still emit exactly its own sheets.
    monkeypatch.setattr(xlsx_module, "Workbook", _NoDefaultSheetWorkbook)
    out = tmp_path / "no-default.xlsx"
    write_workbook(out, [Sheet("Job One", rows())], columns())

    workbook = load_workbook(out)
    try:
        assert workbook.sheetnames == ["Job One"]
        assert workbook["Job One"].cell(row=2, column=1).value == "alpha"
    finally:
        workbook.close()


def test_select_xlsx_columns_internal_has_everything() -> None:
    assert [c.header for c in select_xlsx_columns(columns(), include_costs=True)] == [
        "Name",
        "Amount",
        "Hidden",
    ]


def test_select_xlsx_columns_customer_defaults() -> None:
    selected = select_xlsx_columns(
        columns(), include_costs=False, default_customer_keys=("amount",)
    )
    assert [c.header for c in selected] == ["Name", "Amount"]


def test_select_xlsx_columns_customer_config() -> None:
    selected = select_xlsx_columns(
        columns(), include_costs=False, customer_config={"hidden": True}
    )
    assert [c.header for c in selected] == ["Name", "Hidden"]


def test_select_xlsx_columns_customer_without_config() -> None:
    selected = select_xlsx_columns(columns(), include_costs=False)
    assert [c.header for c in selected] == ["Name"]


def test_write_workbook_round_trip(tmp_path: Path) -> None:
    out = tmp_path / "book.xlsx"
    write_workbook(
        out,
        [Sheet("Job One", rows()), Sheet("Job Two", [])],
        columns(),
        include_costs=True,
    )

    workbook = load_workbook(out)
    assert workbook.sheetnames == ["Job One", "Job Two"]
    sheet = workbook["Job One"]
    assert [cell.value for cell in sheet[1]] == ["Name", "Amount", "Hidden"]
    assert sheet.cell(row=2, column=1).value == "alpha"
    assert sheet.cell(row=2, column=2).value == 1.5
    assert sheet.cell(row=3, column=1).value == "b"
    assert sheet.cell(row=3, column=2).value == 2.0

    header = sheet.cell(row=1, column=1)
    assert header.font.bold is True
    assert header.fill.fill_type == "solid"

    assert workbook["Job Two"].cell(row=1, column=1).value == "Name"


def test_write_workbook_customer_view_drops_disabled_columns(tmp_path: Path) -> None:
    out = tmp_path / "customer.xlsx"
    write_workbook(
        out,
        [Sheet("Only", rows())],
        columns(),
        include_costs=False,
        customer_config={"amount": True},
    )
    sheet = load_workbook(out)["Only"]
    assert [cell.value for cell in sheet[1]] == ["Name", "Amount"]


def test_write_workbook_removes_the_default_sheet(tmp_path: Path) -> None:
    out = tmp_path / "no-default.xlsx"
    write_workbook(out, [Sheet("A", rows())], columns())
    assert "Sheet" not in load_workbook(out).sheetnames


def test_write_workbook_caps_column_width(tmp_path: Path) -> None:
    out = tmp_path / "wide.xlsx"
    long_columns: list[XlsxColumn[tuple[str, float]]] = [
        XlsxColumn("Name", lambda row: row[0])
    ]
    write_workbook(
        out,
        [Sheet("Wide", [("z" * 80, 1.0)])],
        long_columns,
        include_costs=True,
    )
    sheet = load_workbook(out)["Wide"]
    assert sheet.column_dimensions["A"].width == 50


def test_write_workbook_writes_none_cells(tmp_path: Path) -> None:
    out = tmp_path / "none.xlsx"
    nullable: list[XlsxColumn[tuple[str, float]]] = [
        XlsxColumn("Name", lambda row: row[0]),
        XlsxColumn("Blank", lambda row: None),
    ]
    write_workbook(out, [Sheet("S", rows())], nullable, include_costs=True)
    sheet = load_workbook(out)["S"]
    assert sheet.cell(row=2, column=2).value is None


def test_sheet_and_column_dataclasses_are_frozen() -> None:
    sheet: Sheet[tuple[str, float]] = Sheet("n", rows())
    with pytest.raises(Exception):
        sheet.name = "other"  # type: ignore[misc]
    column: XlsxColumn[tuple[str, float]] = XlsxColumn("h", lambda row: row[0])
    assert column.config_key is None


def test_write_workbook_accepts_a_sequence_of_sheets(
    tmp_path: Path,
) -> None:
    out = tmp_path / "seq.xlsx"
    sheets: Sequence[Sheet[tuple[str, float]]] = tuple(
        Sheet(f"S{i}", [row]) for i, row in enumerate(rows())
    )
    factory: Callable[[], list[XlsxColumn[tuple[str, float]]]] = columns
    write_workbook(out, sheets, factory(), include_costs=True)
    assert load_workbook(out).sheetnames == ["S0", "S1"]
