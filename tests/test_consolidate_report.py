"""Tests for :mod:`vinyllabels.consolidate.report`.

The consolidated report is the human-facing deliverable plus its JSON sibling, so
these tests check the headline numbers, the presence (and absence) of each
section, and the per-label text/Excel listings.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from vinyllabels.consolidate.merge import (
    consolidate,
    consolidate_label_sizes,
    consolidate_labels,
    consolidate_size_areas,
)
from vinyllabels.consolidate.model import (
    ConsolidatedMetrics,
    JobReport,
    SizeAreas,
)
from vinyllabels.consolidate.report import (
    LabelRow,
    write_consolidated_report,
    write_consolidated_report_json,
    write_per_label_report,
    write_per_label_report_xlsx,
)

from helpers import cost_breakdown, job_report, label_metrics, pdf_metrics

_SUBRULE = "-" * 80


def normalized(text: str) -> str:
    """Whitespace-collapsed text without column separators.

    The tables are fixed-width with ``|`` between columns; ignoring both makes the
    assertions read like the values they check.
    """
    return " ".join(text.replace("|", " ").split())


def section(text: str, title: str) -> list[str]:
    """The body lines of one report section, excluding its rules."""
    lines = text.splitlines()
    start = lines.index(title)
    body: list[str] = []
    for line in lines[start + 2 :]:
        if not line.strip():
            break
        body.append(line)
    return body


def _job_a(tmp_path: Path) -> JobReport:
    aa_cost = cost_breakdown(1.0, 2.0, 3.0, 4.0, 20.0)
    return job_report(
        input_file=str(tmp_path / "a.txt"),
        labels=[label_metrics("AA", 1.0, aa_cost), label_metrics("BB", 2.0)],
        cost=cost_breakdown(1.0, 2.0, 3.0, 4.0, 20.0),
        pdfs=[tmp_path / "a_2.pdf", tmp_path / "a_1.pdf"],
        per_pdf=[pdf_metrics(2, 4, 3.0, 48.0), pdf_metrics(2, 4, 3.0, 48.0)],
    )


def _job_b(tmp_path: Path) -> JobReport:
    return job_report(
        input_file=str(tmp_path / "b.txt"),
        copies=1,
        size=(4.0, 2.0),
        labels=[label_metrics("AA", 1.0)],
        substrate=24.0,
        text_height_in=3.0,
        page_width_in=24.0,
    )


def _reports(tmp_path: Path) -> list[JobReport]:
    return [_job_a(tmp_path), _job_b(tmp_path)]


def _write(tmp_path: Path, reports: list[JobReport] | None = None) -> tuple[str, Path]:
    """Write a consolidated report, returning its text and its JSON sibling."""
    jobs = _reports(tmp_path) if reports is None else reports
    pricing = None
    metrics = consolidate(jobs, pricing)
    sizes = consolidate_label_sizes(jobs)
    areas = consolidate_size_areas(jobs, pricing)
    labels = consolidate_labels(jobs, pricing)
    out = tmp_path / "Multi-Job Report" / "combined.txt"
    out.parent.mkdir(parents=True, exist_ok=True)
    json_path = write_consolidated_report(
        metrics, sizes, areas, labels, jobs, out, tmp_path / "json"
    )
    return out.read_text(encoding="utf-8"), json_path


# --- LabelRow ------------------------------------------------------------------


def test_label_row_with_a_size() -> None:
    cost = cost_breakdown(1.0, 2.0, 3.0, 4.0, 20.0)
    row = LabelRow("job", label_metrics("AA", 1.0, cost), (8.0, 3.0))
    assert row.label_size == "8x3"
    assert row.label_size_sq_in == 24.0
    assert row.label_area_sq_ft == pytest.approx(24.0 / 144.0)
    assert row.ink_sq_ft == pytest.approx(1.0 / 144.0)
    assert row.breakdown == cost
    assert row.cost("total_cost") == "10.00"
    assert row.cost("unit_price") == "20.00"


def test_label_row_without_a_size_or_pricing() -> None:
    row = LabelRow("job", label_metrics("AA", 1.0))
    assert row.label_size == ""
    assert row.label_size_sq_in == 0.0
    assert row.label_area_sq_ft == 0.0
    assert row.breakdown is None
    # Unpriced rows render blank cells rather than zeros.
    assert row.cost("ink_cost") == ""


# --- JSON report ---------------------------------------------------------------


def test_write_consolidated_report_json(tmp_path: Path) -> None:
    jobs = _reports(tmp_path)
    metrics = consolidate(jobs)
    sizes = consolidate_label_sizes(jobs)
    areas = consolidate_size_areas(jobs)
    labels = consolidate_labels(jobs)
    out = tmp_path / "combined.json"
    write_consolidated_report_json(
        metrics,
        sizes,
        areas,
        labels,
        jobs,
        out,
        tmp_path / "json",
        "2026-01-02 03:04:05",
    )
    report = json.loads(out.read_text(encoding="utf-8"))

    job = report["job"]
    assert isinstance(job, dict)
    # The marker is what keeps a re-run from consolidating its own output.
    assert job["report_type"] == "consolidated"
    assert job["generated"] == "2026-01-02 03:04:05"
    assert job["source_directory"] == str(tmp_path / "json")
    assert job["total_jobs"] == 2
    assert len(job["report_files"]) == 2

    glob = report["global"]
    assert isinstance(glob, dict)
    assert glob["total_output_labels"] == 5
    assert glob["total_substrate_sq_in"] == 120.0
    # with_sq_ft adds a square-foot sibling next to every square-inch field.
    assert glob["total_substrate_sq_ft"] == pytest.approx(120.0 / 144.0)
    assert glob["cost_breakdown"] is None

    sizes_out = report["label_sizes"]
    assert isinstance(sizes_out, list)
    assert [
        (entry["width_in"], entry["height_in"], entry["labels"]) for entry in sizes_out
    ] == [
        (4.0, 2.0, 1),
        (8.0, 3.0, 4),
    ]
    first = sizes_out[0]
    assert first["substrate_sq_in"] == 24.0
    assert first["substrate_sq_ft"] == pytest.approx(24.0 / 144.0)

    jobs_out = report["jobs"]
    assert isinstance(jobs_out, list)
    assert [entry["input_file"] for entry in jobs_out] == [
        str(tmp_path / "a.txt"),
        str(tmp_path / "b.txt"),
    ]
    assert jobs_out[0]["copies_per_label"] == 2
    assert jobs_out[0]["page_width_in"] == 52.0
    assert jobs_out[1]["label_sizes"] == [
        {"width_in": 4.0, "height_in": 2.0, "labels": 1}
    ]

    per_label = report["per_label"]
    assert isinstance(per_label, list)
    assert [entry["text"] for entry in per_label] == ["AA", "BB"]
    assert per_label[0]["instances"] == 3
    assert per_label[0]["ink_area_sq_in"] == 3.0
    assert per_label[0]["ink_area_sq_ft"] == pytest.approx(3.0 / 144.0)
    # label_area_sq_in is a property, so it is not part of the serialised fields.
    assert "label_area_sq_in" not in per_label[0]


# --- Consolidated text report --------------------------------------------------


def test_write_consolidated_report_headline(tmp_path: Path) -> None:
    text, json_path = _write(tmp_path)
    assert json_path == tmp_path / "Multi-Job Report" / "combined.json"
    assert json_path.exists()

    flat = normalized(text)
    assert "VINYL LABEL PRINTING REPORT - CONSOLIDATED" in text
    assert "Jobs Consolidated: 2" in text
    assert (
        "Total Substrate Required: 120.00 sq in ( 0.83 sq ft) (0.24 linear feet of "
        "24in/52in roll)" in flat
    )
    assert "Total Label Area: 104.00 sq in ( 0.72 sq ft)" in flat
    assert "Material Yield: 86.67 %" in flat
    assert "Total Ink Area: 7.00 sq in ( 0.05 sq ft)" in flat
    assert "Average Ink Coverage: 6.73 %" in flat
    assert "Total Character Count: 10" in flat
    assert "Total Output Labels: 5" in flat


def test_write_consolidated_report_omits_costs_without_pricing(tmp_path: Path) -> None:
    text, _ = _write(tmp_path)
    assert "COST BREAKDOWN" not in text


def test_write_consolidated_report_includes_costs_when_priced(
    tmp_path: Path,
) -> None:
    jobs = _reports(tmp_path)
    metrics = consolidate(jobs, None)
    priced = ConsolidatedMetrics(
        total_jobs=metrics.total_jobs,
        total_output_labels=metrics.total_output_labels,
        total_characters=metrics.total_characters,
        total_ink_area_sq_in=metrics.total_ink_area_sq_in,
        total_label_material_sq_in=metrics.total_label_material_sq_in,
        total_substrate_sq_in=metrics.total_substrate_sq_in,
        material_yield_pct=metrics.material_yield_pct,
        average_ink_coverage_pct=metrics.average_ink_coverage_pct,
        linear_feet=metrics.linear_feet,
        cost_breakdown=cost_breakdown(1.0, 2.0, 3.0, 4.0, 20.0),
    )
    out = tmp_path / "combined.txt"
    write_consolidated_report(
        priced,
        consolidate_label_sizes(jobs),
        consolidate_size_areas(jobs),
        consolidate_labels(jobs),
        jobs,
        out,
        tmp_path / "json",
    )
    lines = [
        normalized(line)
        for line in section(out.read_text(encoding="utf-8"), "COST BREAKDOWN")
    ]
    assert lines == [
        "Ink Cost: $ 1.00",
        "Substrate Cost: $ 2.00",
        "Printer Hours: 0.50 hrs",
        "Printer Cost: $ 3.00",
        "Labor Hours: 0.25 hrs",
        "Labor Cost: $ 4.00",
        "Total Cost: $ 10.00",
        "Total Price: $ 20.00",
    ]


def test_write_consolidated_report_size_breakdown(tmp_path: Path) -> None:
    # No label in these jobs carries pricing, so the cost columns are absent
    # entirely; test_write_consolidated_report_size_breakdown_with_costs covers
    # the priced shape of the same table.
    unpriced = [
        job_report(
            input_file=str(tmp_path / "a.txt"),
            labels=[label_metrics("AA", 1.0), label_metrics("BB", 2.0)],
            pdfs=[tmp_path / "a_2.pdf", tmp_path / "a_1.pdf"],
            per_pdf=[pdf_metrics(2, 4, 3.0, 48.0), pdf_metrics(2, 4, 3.0, 48.0)],
        ),
        _job_b(tmp_path),
    ]
    text, _ = _write(tmp_path, unpriced)
    priced = section(text, "LABEL SIZE BREAKDOWN")
    assert normalized(priced[0]) == (
        "Label Size (WxH) Jobs Files Labels Lin Ft Substrate (sq ft) "
        "Label Area (sq ft) Ink (sq ft)"
    )
    assert "Total Cost" not in priced[0]
    assert normalized(priced[2]).startswith("4x2 1 1 1 0.08 0.17 0.06 0.01")
    assert normalized(priced[3]).startswith("8x3 1 2 4 0.15 0.67 0.67 0.04")
    # The totals row reports the consolidated figures, not a re-sum of buckets.
    assert normalized(priced[5]) == "Total 2 3 5 0.24 0.83 0.72 0.05"

    vendor = section(text, "VENDOR LABEL SIZE BREAKDOWN")
    assert normalized(vendor[0]) == normalized(priced[0])


def test_write_consolidated_report_size_breakdown_with_costs(tmp_path: Path) -> None:
    jobs = _reports(tmp_path)
    sizes = consolidate_label_sizes(jobs)
    areas = consolidate_size_areas(jobs)
    priced_areas = {
        size: (
            SizeAreas(
                substrate_sq_in=area.substrate_sq_in,
                label_material_sq_in=area.label_material_sq_in,
                ink_area_sq_in=area.ink_area_sq_in,
                linear_feet=area.linear_feet,
                cost_breakdown=cost_breakdown(1.0, 2.0, 3.0, 4.0, 20.0),
            )
            if size == (8.0, 3.0)
            else area
        )
        for size, area in areas.items()
    }
    metrics = consolidate(jobs)
    out = tmp_path / "combined.txt"
    write_consolidated_report(
        metrics,
        sizes,
        priced_areas,
        consolidate_labels(jobs),
        jobs,
        out,
        tmp_path / "json",
    )
    text = out.read_text(encoding="utf-8")
    priced = section(text, "LABEL SIZE BREAKDOWN")
    assert "Substrate Cost" in priced[0] and "Total Cost" in priced[0]
    # The priced bucket shows money; the unpriced one leaves the cells blank.
    assert normalized(priced[3]).endswith("$ 2.00 $ 1.00 $ 3.00 $ 4.00 $ 10.00")
    assert normalized(priced[2]) == "4x2 1 1 1 0.08 0.17 0.06 0.01"
    # The vendor twin never shows money.
    assert "Total Cost" not in section(text, "VENDOR LABEL SIZE BREAKDOWN")[0]


def test_write_consolidated_report_file_breakdown(tmp_path: Path) -> None:
    text, _ = _write(tmp_path)
    internal = section(text, "FILE BREAKDOWN")
    assert normalized(internal[0]) == (
        "File # Filename Label WxH Labels Lin Ft Substrate (sq ft) "
        "Label Area (sq ft) Ink (sq in) Ink (sq ft)"
    )
    assert normalized(internal[2]) == ("1 a_1.pdf 8x3 2 0.08 0.33 0.33 3.00 0.02")
    assert normalized(internal[3]) == ("2 a_2.pdf 8x3 2 0.08 0.33 0.33 3.00 0.02")
    assert normalized(internal[5]) == "TOTAL 4 0.15 0.67 0.67 6.00 0.04"

    vendor = section(text, "VENDOR FILE BREAKDOWN")
    assert "Filename" not in vendor[0]
    assert normalized(vendor[2]) == "1 8x3 2 0.08 0.33 0.33 3.00 0.02"


def test_write_consolidated_report_skips_file_breakdown_without_records(
    tmp_path: Path,
) -> None:
    jobs = [job_report(input_file="solo.txt", labels=[label_metrics("AA", 1.0)])]
    text, _ = _write(tmp_path, jobs)
    assert "FILE BREAKDOWN" not in text


def test_write_consolidated_report_job_breakdown(tmp_path: Path) -> None:
    text, _ = _write(tmp_path)
    lines = section(text, "JOB BREAKDOWN")
    assert normalized(lines[0]) == (
        "Input File WxH # Files Roll Labels Substrate (sq ft) Label Area (sq ft) "
        "Lin Ft Ink (sq in) Ink (sq ft)"
    )
    # Input paths are shown relative to the report directory's parent.
    assert normalized(lines[2]) == "a.txt 8x3 2 52 4 0.67 0.67 0.15 6.00 0.04"
    assert normalized(lines[3]) == "b.txt 4x2 1 24 1 0.17 0.06 0.08 1.00 0.01"
    assert normalized(lines[5]) == "TOTAL 3 5 0.83 0.72 0.24 7.00 0.05"


def test_write_consolidated_report_per_label_breakdown(tmp_path: Path) -> None:
    text, _ = _write(tmp_path)
    lines = section(text, "PER-LABEL BREAKDOWN")
    assert normalized(lines[0]) == (
        "Label Code Jobs Copies Chars Label Area (sq ft) Substrate (sq ft) "
        "Ink Area (sq in) Ink Area (sq ft)"
    )
    # AA prints twice from job a (8x3) and once from job b (4x2); the row keeps
    # the size it was first seen at, and substrate follows each job's own roll.
    assert normalized(lines[2]) == "AA 2 3 6 0.5000 0.50 3.0000 0.0208"
    assert normalized(lines[3]) == "BB 1 2 4 0.3333 0.33 4.0000 0.0278"
    assert normalized(lines[5]) == "TOTAL 3 5 10 0.8333 0.83 7.0000 0.0486"


def test_write_consolidated_report_honours_an_explicit_json_path(
    tmp_path: Path,
) -> None:
    jobs = _reports(tmp_path)
    metrics = consolidate(jobs)
    out = tmp_path / "combined.txt"
    elsewhere = tmp_path / "elsewhere"
    elsewhere.mkdir()
    json_path = elsewhere / "sibling.json"
    returned = write_consolidated_report(
        metrics,
        consolidate_label_sizes(jobs),
        consolidate_size_areas(jobs),
        consolidate_labels(jobs),
        jobs,
        out,
        tmp_path / "json",
        json_path,
    )
    assert returned == json_path
    assert json_path.exists()
    assert not (tmp_path / "combined.json").exists()


def test_write_consolidated_report_with_a_job_that_recorded_no_sizes(
    tmp_path: Path,
) -> None:
    jobs = [job_report(input_file="weird.txt", size=None, labels=[])]
    text, _ = _write(tmp_path, jobs)
    lines = section(text, "JOB BREAKDOWN")
    # A job that recorded no size renders a blank WxH cell, which collapses out
    # of the normalised row, and still counts as one file on the 52in roll.
    assert lines[2].split("|")[1].strip() == ""
    assert normalized(lines[2]).startswith("weird.txt 1 52 0")


# --- Per-label listings --------------------------------------------------------


def test_write_per_label_report_internal(tmp_path: Path) -> None:
    out = tmp_path / "labels.txt"
    write_per_label_report(_reports(tmp_path), out, tmp_path / "json")
    lines = out.read_text(encoding="utf-8").splitlines()
    # The rules span the rendered table exactly, however wide it grew.
    rule = "-" * len(lines[1])
    assert lines[0] == rule
    assert normalized(lines[1]) == (
        "Job Label Code File Char Count Scale Ink (sq in) Ink (sq ft) Ink ($) "
        "Substrate ($) Print Hrs Print ($) Labor Hrs Labor ($) Total ($) Unit ($)"
    )
    assert lines[2] == rule
    # Job a split into two PDFs: AA's first instance lands in a_2.pdf (the
    # first range) and BB's in a_1.pdf; job b recorded no PDFs, so File is blank.
    assert normalized(lines[3]) == (
        "a AA a_2.pdf 2 1.000 1.0000 0.0069 1.00 2.00 0.50 3.00 0.25 4.00 10.00 20.00"
    )
    # An unpriced label leaves every money cell blank.
    assert normalized(lines[4]) == "a BB a_1.pdf 2 1.000 2.0000 0.0139"
    assert normalized(lines[5]) == "b AA 2 1.000 1.0000 0.0069"
    assert lines[6] == rule
    assert out.read_text(encoding="utf-8").endswith(f"{rule}\n")


def test_write_per_label_report_customer(tmp_path: Path) -> None:
    out = tmp_path / "labels_customer.txt"
    write_per_label_report(
        _reports(tmp_path), out, tmp_path / "json", include_costs=False
    )
    lines = out.read_text(encoding="utf-8").splitlines()
    rule = "-" * len(lines[1])
    assert lines[0] == rule
    assert normalized(lines[1]) == "Job Label Code File Unit Price ($)"
    assert normalized(lines[3]) == "a AA a_2.pdf 20.00"
    assert normalized(lines[4]) == "a BB a_1.pdf"
    assert lines[6] == rule


def test_write_per_label_report_grows_identity_columns_to_fit(tmp_path: Path) -> None:
    long_code = "ABCDEFGHIJKLMNOPQRSTU"
    long_job = "a-much-longer-job-name-than-the-header"
    long_pdf = "some_really_long_output_file_name.pdf"
    jobs = [
        job_report(
            input_file=f"{long_job}.txt",
            labels=[label_metrics(long_code, 1.0)],
            pdfs=[tmp_path / long_pdf],
        )
    ]
    out = tmp_path / "labels.txt"
    write_per_label_report(jobs, out, tmp_path / "json")
    lines = out.read_text(encoding="utf-8").splitlines()
    # Identity columns are never clipped: each grows to the longest value,
    # floored at the header width, so every row prints in full.
    assert normalized(lines[3]).startswith(f"{long_job} {long_code} {long_pdf}")
    cells = lines[3].split(" | ")
    assert len(cells[0]) == len(long_job)
    assert len(cells[1]) == len(long_code)
    assert len(cells[2]) == len(long_pdf)


def test_write_per_label_report_maps_labels_to_the_first_pdf_without_per_pdf(
    tmp_path: Path,
) -> None:
    # A job that recorded PDFs but no per-PDF metrics (a single-PDF run) maps
    # every label to its first, and only, PDF.
    jobs = [
        job_report(
            input_file="a.txt",
            labels=[label_metrics("AA", 1.0), label_metrics("BB", 2.0)],
            pdfs=[tmp_path / "a.pdf"],
        )
    ]
    out = tmp_path / "labels.txt"
    write_per_label_report(jobs, out, tmp_path / "json")
    lines = out.read_text(encoding="utf-8").splitlines()
    assert normalized(lines[3]).startswith("a AA a.pdf 2")
    assert normalized(lines[4]).startswith("a BB a.pdf 2")


# --- Per-label workbooks -------------------------------------------------------


def _sheet_rows(path: Path, sheet: str) -> list[list[object]]:
    """One worksheet as plain rows of cell values, header first."""
    workbook = load_workbook(path)
    try:
        return [list(row) for row in workbook[sheet].iter_rows(values_only=True)]
    finally:
        workbook.close()


def _header(path: Path, sheet: str, row: int = 1) -> list[object]:
    """One row of a worksheet as a list of cell values."""
    return _sheet_rows(path, sheet)[row - 1]


def test_write_per_label_report_xlsx_internal(tmp_path: Path) -> None:
    out = tmp_path / "labels.xlsx"
    write_per_label_report_xlsx(_reports(tmp_path), out, tmp_path / "json")
    assert _header(out, "a") == [
        "Label Code",
        "File",
        "Label Size (WxH)",
        "Char Count",
        "Scale",
        "Ink (sq in)",
        "Ink (sq ft)",
        "Ink ($)",
        "Substrate ($)",
        "Print Hrs",
        "Print ($)",
        "Labor Hrs",
        "Labor ($)",
        "Total ($)",
        "Unit ($)",
    ]
    workbook = load_workbook(out)
    try:
        assert workbook.sheetnames == ["a", "b"]
        assert workbook["a"]["A1"].font.bold is True
    finally:
        workbook.close()
    rows = _sheet_rows(out, "a")
    assert rows[1][:7] == [
        "AA",
        "a_2.pdf",
        "8x3",
        2,
        1.0,
        1.0,
        pytest.approx(1.0 / 144.0),
    ]
    assert rows[1][7:] == [1.0, 2.0, 0.5, 3.0, 0.25, 4.0, 10.0, 20.0]
    # The unpriced label still renders, with blank money cells.
    assert rows[2][7:] == [None] * 8
    # Job b recorded no PDFs, so its File cell is blank (openpyxl reads an
    # empty string back as None).
    assert _sheet_rows(out, "b")[1][:3] == ["AA", None, "4x2"]


def test_write_per_label_report_xlsx_customer_defaults(tmp_path: Path) -> None:
    out = tmp_path / "labels_customer.xlsx"
    write_per_label_report_xlsx(
        _reports(tmp_path), out, tmp_path / "json", include_costs=False
    )
    # Without a config, only the always-on columns plus the unit price show.
    assert _header(out, "a") == ["Label Code", "File", "Unit Price ($)"]


def test_write_per_label_report_xlsx_customer_config(tmp_path: Path) -> None:
    out = tmp_path / "labels_customer.xlsx"
    write_per_label_report_xlsx(
        _reports(tmp_path),
        out,
        tmp_path / "json",
        include_costs=False,
        customer_report_config={"label_size_wxh": True, "char_count": True},
    )
    assert _header(out, "a") == [
        "Label Code",
        "File",
        "Label Size (WxH)",
        "Char Count",
    ]
    assert _sheet_rows(out, "a")[1] == ["AA", "a_2.pdf", "8x3", 2]


def test_write_per_label_report_xlsx_customer_pricing_cells(
    tmp_path: Path,
) -> None:
    out = tmp_path / "labels_customer.xlsx"
    write_per_label_report_xlsx(
        _reports(tmp_path),
        out,
        tmp_path / "json",
        include_costs=False,
        customer_report_config={"total_cost": True, "ink_sq_ft": True},
    )
    rows = _sheet_rows(out, "a")
    assert rows[0] == ["Label Code", "File", "Ink (sq ft)", "Total Cost ($)"]
    assert rows[1] == ["AA", "a_2.pdf", pytest.approx(1.0 / 144.0), 10.0]
    # A row without pricing has no value for the cost column.
    assert rows[2] == ["BB", "a_1.pdf", pytest.approx(2.0 / 144.0), None]


def test_write_per_label_report_xlsx_customer_area_columns(
    tmp_path: Path,
) -> None:
    out = tmp_path / "labels_customer.xlsx"
    write_per_label_report_xlsx(
        _reports(tmp_path),
        out,
        tmp_path / "json",
        include_costs=False,
        customer_report_config={"label_size_sq_in": True, "label_area_sq_ft": True},
    )
    rows = _sheet_rows(out, "a")
    assert rows[0] == [
        "Label Code",
        "File",
        "Label Size (sq in)",
        "Label Area (sq ft)",
    ]
    # The 8x3 design area, and the same figure in square feet.
    assert rows[1] == ["AA", "a_2.pdf", 24.0, pytest.approx(24.0 / 144.0)]


def test_write_per_label_report_xlsx_sanitizes_sheet_names(
    tmp_path: Path,
) -> None:
    jobs = [
        job_report(
            input_file=str(tmp_path / "a*b[c].txt"),
            labels=[label_metrics("AA", 1.0)],
        )
    ]
    out = tmp_path / "labels.xlsx"
    write_per_label_report_xlsx(jobs, out, tmp_path / "json")
    workbook = load_workbook(out)
    try:
        # Each character Excel rejects becomes its own dash; everything else in
        # the stem is kept.
        assert workbook.sheetnames == ["a-b-c-"]
    finally:
        workbook.close()
