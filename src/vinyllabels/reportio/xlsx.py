"""Excel output: one worksheet per job, with styled headers.

Mirrors :mod:`vinyllabels.reportio.table` for the workbook format, so the same
internal/customer column selection drives both CSV and XLSX.
"""

from __future__ import annotations

from collections.abc import Callable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Generic, TypeVar

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

__all__ = ["Sheet", "XlsxColumn", "sanitize_sheet_name", "write_workbook"]

T = TypeVar("T")

Cell = float | int | str | None

# Characters Excel forbids in a sheet name.
_INVALID_SHEET_CHARS = ("/", "\\", "?", "*", ":", "[", "]")
_SHEET_NAME_LIMIT = 31
_MAX_COLUMN_WIDTH = 50

_HEADER_FILL_COLOR = "D3D3D3"


@dataclass(frozen=True)
class XlsxColumn(Generic[T]):
    """One workbook column; ``config_key`` gates it in customer workbooks."""

    header: str
    value: Callable[[T], Cell]
    config_key: str | None = None


@dataclass(frozen=True)
class Sheet(Generic[T]):
    """A named worksheet and the rows to write into it."""

    name: str
    rows: Sequence[T]


def sanitize_sheet_name(name: str) -> str:
    """Return ``name`` made safe and short enough to be an Excel sheet title."""
    sanitized = name
    for char in _INVALID_SHEET_CHARS:
        sanitized = sanitized.replace(char, "-")
    return sanitized[:_SHEET_NAME_LIMIT]


def select_xlsx_columns(
    columns: Sequence[XlsxColumn[T]],
    *,
    include_costs: bool,
    customer_config: dict[str, bool] | None = None,
    default_customer_keys: Sequence[str] = (),
) -> list[XlsxColumn[T]]:
    """Return the workbook columns to emit for the requested audience."""
    if include_costs:
        return list(columns)
    return [
        column
        for column in columns
        if column.config_key is None
        or (
            customer_config.get(column.config_key, False)
            if customer_config is not None
            else column.config_key in default_customer_keys
        )
    ]


def write_workbook(
    output_path: Path,
    sheets: Sequence[Sheet[T]],
    columns: Sequence[XlsxColumn[T]],
    *,
    include_costs: bool = True,
    customer_config: dict[str, bool] | None = None,
    default_customer_keys: Sequence[str] = (),
) -> None:
    """Write one worksheet per sheet in ``sheets``, styled with bold headers.

    Column widths are auto-fitted to the widest cell, capped so a long label
    code cannot blow out the layout.
    """
    selected = select_xlsx_columns(
        columns,
        include_costs=include_costs,
        customer_config=customer_config,
        default_customer_keys=default_customer_keys,
    )

    workbook = Workbook()
    default_sheet = workbook.active
    if default_sheet is not None:
        workbook.remove(default_sheet)

    header_fill = PatternFill(
        start_color=_HEADER_FILL_COLOR, end_color=_HEADER_FILL_COLOR, fill_type="solid"
    )
    header_font = Font(bold=True)

    for sheet in sheets:
        worksheet = workbook.create_sheet(title=sanitize_sheet_name(sheet.name))

        for col_num, column in enumerate(selected, start=1):
            cell = worksheet.cell(row=1, column=col_num, value=column.header)
            cell.fill = header_fill
            cell.font = header_font

        row_num = 2
        for row in sheet.rows:
            for col_num, column in enumerate(selected, start=1):
                worksheet.cell(row=row_num, column=col_num, value=column.value(row))
            row_num += 1

        for col_num in range(1, len(selected) + 1):
            widest = max(
                len(str(worksheet.cell(row=r, column=col_num).value or ""))
                for r in range(1, row_num)
            )
            letter = worksheet.cell(row=1, column=col_num).column_letter
            worksheet.column_dimensions[letter].width = min(widest, _MAX_COLUMN_WIDTH)

    workbook.save(output_path)
